"""Audit independent molecular/plate evidence; never produce RNA editing truth."""
from pathlib import Path
import re,json,shutil
import numpy as np
import openpyxl
from Bio import SeqIO, Align
from Bio.Seq import Seq
from inventory import ROOT,WORKSPACE,sha,write_csv,dump
from sample_metadata import molecular_metadata

def run():
    src=WORKSPACE/'Reporter/序列/REPORTER_assembled_from_4_primers.gb'
    dest=ROOT/'inputs/construct_reference'/src.name;dest.parent.mkdir(exist_ok=True)
    if not dest.exists():shutil.copy2(src,dest)
    if src.exists():assert sha(src)==sha(dest)
    rec=SeqIO.read(dest,'genbank');seq=str(rec.seq).upper()
    features=[dict(source=str(dest.relative_to(ROOT)),label=';'.join(f.qualifiers.get('label',[])),start_1based=int(f.location.start)+1,end_1based=int(f.location.end),strand=f.location.strand,type=f.type,role='general_design_only_exact_S7_S10_clone_unverified') for f in rec.features]
    write_csv(ROOT/'inventory/construct_features.csv',features)
    wb=openpyxl.load_workbook(ROOT/'inputs/S10/523RepEdit.xlsx',data_only=True);s=wb['Sheet1']
    forward=s['C13'].value.strip(); reverse=s['C14'].value.strip();a=seq.find(forward);b=seq.find(str(Seq(reverse).reverse_complement()));ref=seq[a:b+len(reverse)]
    dump(ROOT/'inventory/primer_reference_audit.json',dict(primer_source='inputs/S10/523RepEdit.xlsx#Sheet1!C13:C14',forward=forward,reverse=reverse,general_design_sha256=sha(dest),general_design_length=len(seq),forward_start_1based=a+1,reverse_binding_start_1based=b+1,design_amplicon_length=len(ref),primer_name_length_label=523,exact_S10_OpenA_reference_verified=False,interpretation='Primers delimit a reporter region in the general design, but the 544 bp design interval is not the exact 523-labelled OpenA amplicon. No molecular target editing percentage is inferred.'))
    aligner=Align.PairwiseAligner(mode='local',match_score=2,mismatch_score=-3,open_gap_score=-6,extend_gap_score=-.3)
    sanger=[];qtable=[]
    for p in sorted((ROOT/'inputs/S10/测序').rglob('*.ab1')):
        r=SeqIO.read(p,'abi');raw=r.annotations['abif_raw'];q=np.array(r.letter_annotations['phred_quality']);label=p.name.split('.')[0];direction=p.name.split('.')[1]
        ss=str(r.seq);ori=ss if direction=='F' else str(r.seq.reverse_complement());al=aligner.align(ref,ori)[0];idx=al.indices;valid=(idx[0]>=0)&(idx[1]>=0)
        identical=sum(ref[i]==ori[j] for i,j in idx[:,valid].T)
        order=raw['FWO_1'].decode();traces={base:np.array(raw[f'DATA{i+9}']) for i,base in enumerate(order)}
        peaks=np.array(raw['PLOC2']);signals=np.stack([traces[base][peaks] for base in 'ACGT'],axis=1);sort=np.sort(signals,axis=1);second=np.divide(sort[:,-2],signals.sum(axis=1),out=np.zeros(len(q)),where=signals.sum(axis=1)>0)
        interior=(np.arange(len(q))>=30)&(np.arange(len(q))<len(q)-30)&(q>=20)
        source_seq=p.with_suffix('.seq');seqtext=source_seq.read_text(errors='replace') if source_seq.exists() else ''
        called=''.join(x.strip() for x in seqtext.splitlines() if not x.startswith('>')).upper()
        sanger.append(dict(read_id=p.stem,source=str(p.relative_to(ROOT)),sha256=sha(p),sample_label=label,direction=direction,called_bases=len(r),Q20_bases=int((q>=20).sum()),Q30_bases=int((q>=30).sum()),median_phred=float(np.median(q)),run_date=raw.get('RUND1'),machine=raw.get('MCHN1',b'').decode(errors='replace'),reference_alignment_paired_bases=int(valid.sum()),reference_aligned_identity=identical/int(valid.sum()),reference_type='general_reporter_design_not_exact_OpenA',seq_export_matches_AB1=called==ss,interior_Q20_n=int(interior.sum()),interior_median_second_peak_share=float(np.median(second[interior])),vendor_double_peak_flag=label=='1-8-2' and direction=='F',sample_to_flow_pair=None,editing_fraction=None,editing_status='NOT_ESTIMABLE_ORIGIN_SITE_PAIRING_UNVERIFIED'))
        pairing=molecular_metadata(label)
        sanger[-1].update(pairing,sample_to_flow_pair=pairing['flow_acquisition_pair'],editing_status='NOT_ESTIMABLE_EXACT_SITE_AND_ASSAY_UNQUALIFIED')
        qtable.extend(dict(read_id=p.stem,sample_label=label,direction=direction,base_index_1based=int(i+1),phred=int(q[i]),secondary_peak_share=float(second[i])) for i in range(len(q)))
    write_csv(ROOT/'inventory/molecular_reads.csv',sanger);write_csv(ROOT/'results/sanger_base_quality.csv',qtable)
    pairs=[]
    for label in sorted(set(r['sample_label'] for r in sanger)):
        subset=[r for r in sanger if r['sample_label']==label]
        pairs.append(dict(molecular_sample_label=label,read_n=len(subset),directions=';'.join(r['direction'] for r in subset),independent_biological_culture_id=None,extraction_time=None,RNA_origin='RNA_extracted_after_flow_user_confirmed',DNase_noRT_control=None,endpoint='reporter_region_candidate_not_endogenous_APOE_proof',PER_status='NOT_CALIBRATED',**molecular_metadata(label)))
    write_csv(ROOT/'inventory/molecular_pairing_registry.csv',pairs)
    import pandas as pd
    flow=pd.read_csv(ROOT/'results/ras_results.csv')
    well_rows=[]
    for aid, subset in pd.DataFrame(pairs).groupby('flow_acquisition_pair'):
        fr=flow[(flow.acquisition_id==aid)&(flow.population_mode=='reporter')].iloc[0]
        br=flow[(flow.acquisition_id==aid)&(flow.population_mode=='blue_positive_reporter')].iloc[0]
        well_rows.append(dict(flow_acquisition_id=aid,well_id=fr.well_id,condition_label=fr.condition_label,replicate_label=fr.replicate_label,reporter_dna_ng=fr.reporter_dna_ng,editor_dna_ng=fr.editor_dna_ng,total_dna_ng=fr.total_dna_ng,molecular_labels=';'.join(subset.molecular_sample_label),molecular_label_n=len(subset),read_n=int(subset.read_n.sum()),independent_transfection_well_n=1,RAS_reporter=fr.ras,RAS_editor_coexpression=br.ras if fr.condition_label>0 else None,CEF_editor_exploratory=fr.CEF_editor_exploratory,flow_RNA_pairing_status='USER_CONFIRMED_SAME_WELL_FLOW_THEN_RNA',RNA_sampling_relation=subset.RNA_sampling_relation.iloc[0],RNA_population_selection=subset.RNA_population_selection.iloc[0],RNA_editing_fraction=None,molecular_endpoint_status='NOT_QUALIFIED_EXACT_REFERENCE_SITE_AND_ASSAY_REQUIRED',PER_status='NOT_CALIBRATED'))
    write_csv(ROOT/'results/flow_RNA_well_registry.csv',well_rows)
    # Only the EDIT block of the shared plate map is authorized as S10.
    mapwb=openpyxl.load_workbook(ROOT/'inputs/S10/酶标仪/孔图备注.xlsx',data_only=True)['Sheet1']
    mapping={f'{chr(65+r-13)}{c-1:02d}':mapwb.cell(r,c).value for r in range(13,16) for c in range(2,6)}
    plate=[];settings=[]
    folder=ROOT/'inputs/S10/酶标仪'
    for p in sorted(folder.glob('*.xlsx')):
        if p.name=='孔图备注.xlsx':continue
        w=openpyxl.load_workbook(p,data_only=True);blank='blank' in p.name;multi=w['荧光 1_01'].max_row>30
        for ch,sheet in [('BFP','荧光 1_01'),('EGFP','荧光 2_02'),('mCherry','荧光 3_03')]:
            s=w[sheet];settings.append(dict(file=str(p.relative_to(ROOT)),channel=ch,wavelength=s['A6'].value,run_time=s['A3'].value,spatial_scan=multi,instrument='Varioskan LUX',optical_response_compensation='No',blank_wells_documented='A01 only' if blank else None))
            if multi:
                for row in s:
                    for cell in row:
                        if not isinstance(cell.value,str) or not re.fullmatch('[A-C]0[1-4]',cell.value):continue
                        well=cell.value
                        if blank and well!='A01':continue
                        for rr in range(cell.row+1,cell.row+12):
                            for cc in range(cell.column+1,cell.column+12):
                                val=s.cell(rr,cc).value
                                if isinstance(val,(float,int)):
                                    plate.append(dict(file=p.name,sheet=sheet,source_cell=s.cell(rr,cc).coordinate,well=well,sample_label='PBS' if blank else mapping[well],channel=ch,x=s.cell(cell.row,cc).value,y=s.cell(rr,cell.column).value,value_RFU=val,kind='PBS_spatial_scan' if blank else 'S10_spatial_scan',technical_point_not_replicate=True))
            else:
                for rr in range(11,14):
                    for cc in range(2,6):
                        well=f'{chr(65+rr-11)}{cc-1:02d}'
                        plate.append(dict(file=p.name,sheet=sheet,source_cell=s.cell(rr,cc).coordinate,well=well,sample_label=mapping[well],channel=ch,x=None,y=None,value_RFU=s.cell(rr,cc).value,kind='S10_single_point',technical_point_not_replicate=True))
    import pandas as pd
    d=pd.DataFrame(plate);summ=d.groupby(['file','kind','well','sample_label','channel']).agg(point_n=('value_RFU','count'),mean_RFU=('value_RFU','mean'),sd_spatial=('value_RFU','std'),minimum=('value_RFU','min'),maximum=('value_RFU','max')).reset_index()
    blanks=summ[summ.kind=='PBS_spatial_scan'].set_index('channel').mean_RFU.to_dict()
    summ['diagnostic_mean_minus_PBS_scan_mean']=[v-blanks[ch] for v,ch in zip(summ.mean_RFU,summ.channel)]
    summ['blank_compatibility']='different_time_same_scan_only_single_point_not_matched'
    write_csv(ROOT/'results/plate_raw_points.csv',plate);write_csv(ROOT/'results/plate_summary.csv',summ.to_dict('records'));write_csv(ROOT/'inventory/plate_settings.csv',settings)
    print('Molecular reads',len(sanger),'sample labels',len(pairs),'plate points retained',len(plate))

if __name__=='__main__':run()
